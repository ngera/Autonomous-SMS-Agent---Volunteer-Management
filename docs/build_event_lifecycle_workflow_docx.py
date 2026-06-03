"""Generate event_lifecycle_workflow.docx from the markdown source.

Run: python docs/build_event_lifecycle_workflow_docx.py
Input:  docs/event_lifecycle_workflow.md
Output: docs/event_lifecycle_workflow.docx

Handles: H1/H2/H3 headings, paragraphs, bullet + numbered lists, GFM tables,
fenced code blocks, horizontal rules, inline **bold**, *italic*, and `code`.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor


INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
)


def add_inline(paragraph, text: str) -> None:
    """Append `text` to `paragraph`, honouring **bold**, *italic*, `code`."""
    pos = 0
    for m in INLINE_PATTERN.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_cell_shading(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def parse_table(lines: list[str], i: int) -> tuple[list[list[str]], int]:
    """Consume markdown table starting at `i`; return (rows, next_index)."""
    rows: list[list[str]] = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        raw = lines[i].strip()
        # Strip leading/trailing pipes
        if raw.startswith("|"):
            raw = raw[1:]
        if raw.endswith("|"):
            raw = raw[:-1]
        cells = [c.strip() for c in raw.split("|")]
        # Skip alignment row (| --- | --- |)
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            add_inline(p, text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
                set_cell_shading(cell, "E8EEF7")
    # spacer paragraph below table
    doc.add_paragraph("")


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    run = p.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_list_item(doc: Document, text: str, ordered: bool) -> None:
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    for run in p.runs:
        run.font.size = Pt(11)


def heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(level=level)
    add_inline(h, text)


def convert(md_path: Path, docx_path: Path) -> None:
    raw = md_path.read_text(encoding="utf-8")
    lines = raw.splitlines()

    doc = Document()

    # Set base paragraph style font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Fenced code blocks
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            add_horizontal_rule(doc)
            i += 1
            continue

        # Tables
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # Headings
        if stripped.startswith("# "):
            heading(doc, stripped[2:].strip(), level=0)
            # subtitle paragraph
            sub_idx = i + 1
            while sub_idx < len(lines) and not lines[sub_idx].strip():
                sub_idx += 1
            if sub_idx < len(lines) and lines[sub_idx].startswith("*") and lines[sub_idx].endswith("*"):
                sub_text = lines[sub_idx].strip("*").strip()
                p = doc.add_paragraph()
                run = p.add_run(sub_text)
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                i = sub_idx + 1
                continue
            i += 1
            continue
        if stripped.startswith("## "):
            heading(doc, stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            heading(doc, stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("#### "):
            heading(doc, stripped[5:].strip(), level=3)
            i += 1
            continue

        # Numbered list
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            add_list_item(doc, num_match.group(2), ordered=True)
            i += 1
            continue

        # Bulleted list
        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:], ordered=False)
            i += 1
            continue

        # Plain paragraph — concatenate consecutive non-blank, non-special lines
        para_lines = [stripped]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if not nxt:
                break
            if (
                nxt.startswith("#")
                or nxt.startswith("- ")
                or nxt.startswith("|")
                or nxt.startswith("```")
                or nxt == "---"
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            para_lines.append(nxt)
            j += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(para_lines))
        i = j

    doc.save(docx_path)
    print(f"Wrote {docx_path}")


def main() -> None:
    here = Path(__file__).parent
    convert(here / "event_lifecycle_workflow.md", here / "event_lifecycle_workflow.docx")


if __name__ == "__main__":
    main()
